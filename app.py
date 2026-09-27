import streamlit as st
import pandas as pd
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURAÇÃO
# ============================================================

st.set_page_config(
    page_title="VALORES 6",
    page_icon="💰",
    layout="wide"
)


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

MESES = [
    "Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
    "Jul", "Ago", "Set", "Out", "Nov", "Dez"
]


def dinheiro(valor):
    return (
        f"R$ {valor:,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def nome_mes(ano, mes):
    return f"{MESES[mes - 1]}/{str(ano)[2:]}"


def gerar_meses(ano_inicio, mes_inicio, ano_fim, mes_fim):

    meses = []

    ano = ano_inicio
    mes = mes_inicio

    while (ano < ano_fim) or (ano == ano_fim and mes <= mes_fim):

        meses.append((ano, mes))

        mes += 1

        if mes > 12:
            mes = 1
            ano += 1

    return meses


# ============================================================
# LOGIN
# ============================================================

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if "resultados" not in st.session_state:
    st.session_state.resultados = None


if not st.session_state.autenticado:

    st.title("💰 VALORES 6")

    st.subheader("Acesso ao sistema")

    usuario = st.text_input("Usuário")

    senha = st.text_input(
        "Senha",
        type="password"
    )

    entrar = st.button(
        "🔐 ENTRAR",
        use_container_width=True
    )

    if entrar:

        if usuario == "admin" and senha == "1234":

            st.session_state.autenticado = True
            st.rerun()

        else:

            st.error("Usuário ou senha inválidos.")

    st.stop()


# ============================================================
# TÍTULO
# ============================================================

st.title("💰 VALORES 6")

st.caption(
    "Simulador de projeção financeira"
)


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.header("⚙️ PARÂMETROS")


# ============================================================
# PERÍODO
# ============================================================

st.sidebar.subheader("📅 Período")

ano_inicio = st.sidebar.number_input(
    "Ano inicial",
    min_value=2000,
    max_value=2100,
    value=2026,
    step=1
)

mes_inicio = st.sidebar.selectbox(
    "Mês inicial",
    range(1, 13),
    index=8,
    format_func=lambda x: MESES[x - 1]
)

ano_fim = st.sidebar.number_input(
    "Ano final",
    min_value=2000,
    max_value=2100,
    value=2032,
    step=1
)

mes_fim = st.sidebar.selectbox(
    "Mês final",
    range(1, 13),
    index=11,
    format_func=lambda x: MESES[x - 1]
)

inflacao = st.sidebar.number_input(
    "Inflação mensal (%)",
    min_value=0.0,
    value=0.40,
    step=0.01,
    format="%.2f"
)


# ============================================================
# VALORES INICIAIS
# ============================================================

st.sidebar.subheader("💰 Valores iniciais")

st.sidebar.info(
    "Digite seus valores apenas nesta tela. "
    "Eles não ficam gravados no código-fonte nem no GitHub."
)


salario = st.sidebar.number_input(
    "Salário inicial (R$)",
    min_value=0.0,
    value=None,
    step=100.0,
    format="%.2f",
    placeholder="Digite o salário"
)


prev_saldo = st.sidebar.number_input(
    "Saldo inicial Previcel (R$)",
    min_value=0.0,
    value=None,
    step=1000.0,
    format="%.2f",
    placeholder="Digite o saldo"
)


fgts_saldo = st.sidebar.number_input(
    "Saldo inicial FGTS (R$)",
    min_value=0.0,
    value=None,
    step=1000.0,
    format="%.2f",
    placeholder="Digite o saldo"
)


rf_saldo = st.sidebar.number_input(
    "Saldo inicial Renda Fixa (R$)",
    min_value=0.0,
    value=None,
    step=1000.0,
    format="%.2f",
    placeholder="Digite o saldo"
)


tesouro_saldo = st.sidebar.number_input(
    "Saldo inicial Tesouro (R$)",
    min_value=0.0,
    value=None,
    step=1000.0,
    format="%.2f",
    placeholder="Digite o saldo"
)


# ============================================================
# PREVICEL
# ============================================================

st.sidebar.subheader("🏦 Previcel")

prev_percentual = st.sidebar.number_input(
    "Contribuição participante (%)",
    min_value=0.0,
    value=10.50,
    step=0.10,
    format="%.2f"
)

patro_percentual = st.sidebar.number_input(
    "Contrapartida empresa (%)",
    min_value=0.0,
    value=100.00,
    step=1.00,
    format="%.2f"
)

extra = st.sidebar.number_input(
    "Aporte extra mensal (R$)",
    min_value=0.0,
    value=1050.00,
    step=50.0,
    format="%.2f"
)

aumento_extra = st.sidebar.number_input(
    "Aumento anual do aporte extra (R$)",
    min_value=0.0,
    value=100.00,
    step=10.0,
    format="%.2f"
)

reajuste = st.sidebar.number_input(
    "Reajuste anual do salário (%)",
    min_value=0.0,
    value=4.00,
    step=0.10,
    format="%.2f"
)

rent_prev = st.sidebar.number_input(
    "Rentabilidade mensal Previcel (%)",
    min_value=0.0,
    value=1.00,
    step=0.01,
    format="%.2f"
)

decimo = st.sidebar.checkbox(
    "Contribuir no 13º salário",
    value=True
)


# ============================================================
# FGTS
# ============================================================

st.sidebar.subheader("🏠 FGTS")

fgts_percentual = st.sidebar.number_input(
    "Depósito mensal FGTS (%)",
    min_value=0.0,
    value=8.00,
    step=0.10,
    format="%.2f"
)

fgts_saque = st.sidebar.number_input(
    "Saque anual (R$)",
    min_value=0.0,
    value=6000.00,
    step=500.0,
    format="%.2f"
)

fgts_mes_saque = st.sidebar.selectbox(
    "Mês do saque",
    range(1, 13),
    index=4,
    format_func=lambda x: MESES[x - 1]
)

rent_fgts = st.sidebar.number_input(
    "Rentabilidade mensal FGTS (%)",
    min_value=0.0,
    value=0.30,
    step=0.01,
    format="%.2f"
)


# ============================================================
# RENDA FIXA
# ============================================================

st.sidebar.subheader("💵 Renda Fixa")

rf_aporte = st.sidebar.number_input(
    "Aporte mensal (R$)",
    min_value=0.0,
    value=2300.00,
    step=100.0,
    format="%.2f"
)

rf_aumento = st.sidebar.number_input(
    "Aumento anual do aporte (R$)",
    min_value=0.0,
    value=150.00,
    step=50.0,
    format="%.2f"
)

rf_extra_maio = st.sidebar.number_input(
    "Aporte extra em maio (R$)",
    min_value=0.0,
    value=18000.00,
    step=500.0,
    format="%.2f"
)

rent_rf = st.sidebar.number_input(
    "Rentabilidade mensal Renda Fixa (%)",
    min_value=0.0,
    value=0.70,
    step=0.01,
    format="%.2f"
)


# ============================================================
# TESOURO
# ============================================================

st.sidebar.subheader("📈 Tesouro Direto")

tesouro_aporte = st.sidebar.number_input(
    "Aporte mensal (R$)",
    min_value=0.0,
    value=0.00,
    step=100.0,
    format="%.2f"
)

rent_tesouro = st.sidebar.number_input(
    "Rentabilidade mensal Tesouro (%)",
    min_value=0.0,
    value=0.85,
    step=0.01,
    format="%.2f"
)


# ============================================================
# BOTÕES
# ============================================================

st.sidebar.markdown("---")

calcular = st.sidebar.button(
    "📊 CALCULAR PROJEÇÃO",
    use_container_width=True
)

limpar = st.sidebar.button(
    "🧹 LIMPAR",
    use_container_width=True
)

sair = st.sidebar.button(
    "🚪 SAIR",
    use_container_width=True
)


# ============================================================
# LIMPAR
# ============================================================

if limpar:

    st.session_state.resultados = None

    st.rerun()


# ============================================================
# SAIR
# ============================================================

if sair:

    st.session_state.autenticado = False
    st.session_state.resultados = None

    st.rerun()


# ============================================================
# VALIDAÇÃO E CÁLCULO
# ============================================================

if calcular:

    campos_iniciais = [
        salario,
        prev_saldo,
        fgts_saldo,
        rf_saldo,
        tesouro_saldo
    ]

    if any(valor is None for valor in campos_iniciais):

        st.error(
            "⚠️ Preencha os cinco valores iniciais antes de calcular."
        )

    elif ano_inicio > ano_fim:

        st.error(
            "⚠️ O ano inicial não pode ser maior que o ano final."
        )

    elif ano_inicio == ano_fim and mes_inicio > mes_fim:

        st.error(
            "⚠️ O mês inicial não pode ser posterior ao mês final."
        )

    else:

        meses = gerar_meses(
            ano_inicio,
            mes_inicio,
            ano_fim,
            mes_fim
        )

        salario_atual = float(salario)

        prev_atual = float(prev_saldo)
        fgts_atual = float(fgts_saldo)
        rf_atual = float(rf_saldo)
        tesouro_atual = float(tesouro_saldo)

        extra_atual = float(extra)
        rf_aporte_atual = float(rf_aporte)

        resultados = []

        for indice, (ano, mes) in enumerate(meses):

            # ------------------------------------------------
            # REAJUSTES ANUAIS
            # ------------------------------------------------

            if ano > ano_inicio:

                if mes == 1:

                    extra_atual += aumento_extra
                    rf_aporte_atual += rf_aumento

                if mes == 5:

                    salario_atual *= (
                        1 + reajuste / 100
                    )

            # ------------------------------------------------
            # PREVICEL
            # ------------------------------------------------

            participante = (
                salario_atual *
                prev_percentual / 100
            )

            empresa = (
                participante *
                patro_percentual / 100
            )

            rendimento_prev = (
                prev_atual *
                rent_prev / 100
            )

            decimo_valor = 0.0

            if decimo and mes == 12:

                decimo_valor = (
                    participante + empresa
                )

            prev_atual += (
                rendimento_prev
                + participante
                + empresa
                + extra_atual
                + decimo_valor
            )

            # ------------------------------------------------
            # FGTS
            # ------------------------------------------------

            rendimento_fgts = (
                fgts_atual *
                rent_fgts / 100
            )

            deposito_fgts = (
                salario_atual *
                fgts_percentual / 100
            )

            saque = 0.0

            if mes == fgts_mes_saque and ano > ano_inicio:

                saque = min(
                    fgts_saque,
                    fgts_atual
                    + rendimento_fgts
                    + deposito_fgts
                )

            fgts_atual = (
                fgts_atual
                + rendimento_fgts
                + deposito_fgts
                - saque
            )

            # ------------------------------------------------
            # RENDA FIXA
            # ------------------------------------------------

            rendimento_rf = (
                rf_atual *
                rent_rf / 100
            )

            extra_maio = 0.0

            if mes == 5 and ano > ano_inicio:

                extra_maio = (
                    rf_extra_maio *
                    (1 + reajuste / 100) ** (ano - ano_inicio - 1)
                )

            rf_atual += (
                rendimento_rf
                + rf_aporte_atual
                + extra_maio
            )

            # ------------------------------------------------
            # TESOURO
            # ------------------------------------------------

            rendimento_tesouro = (
                tesouro_atual *
                rent_tesouro / 100
            )

            tesouro_atual += (
                rendimento_tesouro
                + tesouro_aporte
            )

            # ------------------------------------------------
            # TOTAL
            # ------------------------------------------------

            total = (
                prev_atual
                + fgts_atual
                + rf_atual
                + tesouro_atual
            )

            valor_hoje = (
                total /
                (
                    (1 + inflacao / 100)
                    ** (indice + 1)
                )
            )

            resultados.append(
                {
                    "Mês": nome_mes(ano, mes),
                    "Ano": ano,
                    "Salário": salario_atual,
                    "Participante": participante,
                    "Empresa": empresa,
                    "Extra Previcel": extra_atual,
                    "13º": decimo_valor,
                    "Previcel": prev_atual,
                    "FGTS": fgts_atual,
                    "Renda Fixa": rf_atual,
                    "Tesouro": tesouro_atual,
                    "Total": total,
                    "Dinheiro de hoje": valor_hoje
                }
            )

        st.session_state.resultados = pd.DataFrame(
            resultados
        )


# ============================================================
# RESULTADOS
# ============================================================

if st.session_state.resultados is not None:

    df = st.session_state.resultados

    ultimo = df.iloc[-1]

    st.markdown("---")

    st.subheader("📊 Resultado da projeção")

    col1, col2, col3, col4, col5 = st.columns(5)

    col1.metric(
        "Previcel",
        dinheiro(ultimo["Previcel"])
    )

    col2.metric(
        "FGTS",
        dinheiro(ultimo["FGTS"])
    )

    col3.metric(
        "Renda Fixa",
        dinheiro(ultimo["Renda Fixa"])
    )

    col4.metric(
        "Tesouro",
        dinheiro(ultimo["Tesouro"])
    )

    col5.metric(
        "TOTAL",
        dinheiro(ultimo["Total"])
    )


    # ========================================================
    # TABELA MENSAL
    # ========================================================

    st.subheader("📅 Evolução mensal")

    df_exibicao = df.copy()

    colunas_dinheiro = [
        "Salário",
        "Participante",
        "Empresa",
        "Extra Previcel",
        "13º",
        "Previcel",
        "FGTS",
        "Renda Fixa",
        "Tesouro",
        "Total",
        "Dinheiro de hoje"
    ]

    for coluna in colunas_dinheiro:

        df_exibicao[coluna] = df_exibicao[
            coluna
        ].map(dinheiro)

    st.dataframe(
        df_exibicao,
        use_container_width=True,
        hide_index=True
    )


    # ========================================================
    # GRÁFICO
    # ========================================================

    st.subheader("📈 Evolução dos investimentos")

    grafico = df[
        [
            "Mês",
            "Previcel",
            "FGTS",
            "Renda Fixa",
            "Tesouro"
        ]
    ].copy()

    grafico = grafico.set_index("Mês")

    st.line_chart(
        grafico,
        use_container_width=True
    )


    # ========================================================
    # RESUMO ANUAL
    # ========================================================

    st.subheader("📆 Resumo anual")

    resumo_anual = (
        df.groupby("Ano")
        .agg(
            {
                "Previcel": "last",
                "FGTS": "last",
                "Renda Fixa": "last",
                "Tesouro": "last",
                "Total": "last",
                "Dinheiro de hoje": "last"
            }
        )
        .reset_index()
    )

    resumo_exibicao = resumo_anual.copy()

    for coluna in [
        "Previcel",
        "FGTS",
        "Renda Fixa",
        "Tesouro",
        "Total",
        "Dinheiro de hoje"
    ]:

        resumo_exibicao[coluna] = (
            resumo_exibicao[coluna]
            .map(dinheiro)
        )

    st.dataframe(
        resumo_exibicao,
        use_container_width=True,
        hide_index=True
    )


    # ========================================================
    # EXPORTAÇÃO EXCEL
    # ========================================================

    def gerar_excel(df, resumo_anual):

        arquivo = BytesIO()

        wb = Workbook()

        ws_resumo = wb.active
        ws_resumo.title = "Resumo"

        ws_mensal = wb.create_sheet(
            "Evolução Mensal"
        )

        ws_anual = wb.create_sheet(
            "Resumo Anual"
        )

        ws_graficos = wb.create_sheet(
            "Gráficos"
        )

        # ----------------------------------------------------
        # ESTILOS
        # ----------------------------------------------------

        titulo_fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        cabecalho_fill = PatternFill(
            "solid",
            fgColor="5B9BD5"
        )

        branco = Font(
            color="FFFFFF",
            bold=True
        )

        negrito = Font(
            bold=True
        )

        alinhamento = Alignment(
            horizontal="center",
            vertical="center"
        )

        borda = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # ----------------------------------------------------
        # RESUMO
        # ----------------------------------------------------

        ws_resumo["A1"] = "VALORES 6 - RESUMO"

        ws_resumo["A1"].font = Font(
            bold=True,
            size=16,
            color="FFFFFF"
        )

        ws_resumo["A1"].fill = titulo_fill

        ws_resumo.merge_cells(
            "A1:B1"
        )

        dados_resumo = [
            ["Indicador", "Valor"],
            ["Previcel", ultimo["Previcel"]],
            ["FGTS", ultimo["FGTS"]],
            ["Renda Fixa", ultimo["Renda Fixa"]],
            ["Tesouro", ultimo["Tesouro"]],
            ["Total", ultimo["Total"]],
            [
                "Dinheiro de hoje",
                ultimo["Dinheiro de hoje"]
            ]
        ]

        for linha in dados_resumo:

            ws_resumo.append(linha)

        for cell in ws_resumo[2]:

            cell.font = branco
            cell.fill = cabecalho_fill
            cell.alignment = alinhamento
            cell.border = borda

        for row in ws_resumo.iter_rows(
            min_row=3,
            max_row=ws_resumo.max_row
        ):

            for cell in row:

                cell.border = borda

                if cell.column == 2:

                    cell.number_format = (
                        'R$ #,##0.00'
                    )

        ws_resumo.column_dimensions["A"].width = 25
        ws_resumo.column_dimensions["B"].width = 20


        # ----------------------------------------------------
        # EVOLUÇÃO MENSAL
        # ----------------------------------------------------

        for coluna in df.columns:

            ws_mensal.cell(
                row=1,
                column=list(df.columns).index(coluna) + 1,
                value=coluna
            )

        for cell in ws_mensal[1]:

            cell.font = branco
            cell.fill = cabecalho_fill
            cell.alignment = alinhamento
            cell.border = borda

        for r_idx, row in enumerate(
            df.itertuples(index=False),
            start=2
        ):

            for c_idx, valor in enumerate(
                row,
                start=1
            ):

                cell = ws_mensal.cell(
                    row=r_idx,
                    column=c_idx,
                    value=valor
                )

                cell.border = borda

                if (
                    df.columns[c_idx - 1]
                    in colunas_dinheiro
                ):

                    cell.number_format = (
                        'R$ #,##0.00'
                    )

        for coluna in range(
            1,
            ws_mensal.max_column + 1
        ):

            ws_mensal.column_dimensions[
                chr(64 + coluna)
            ].width = 18


        # ----------------------------------------------------
        # RESUMO ANUAL
        # ----------------------------------------------------

        for coluna in resumo_anual.columns:

            ws_anual.cell(
                row=1,
                column=list(
                    resumo_anual.columns
                ).index(coluna) + 1,
                value=coluna
            )

        for cell in ws_anual[1]:

            cell.font = branco
            cell.fill = cabecalho_fill
            cell.alignment = alinhamento
            cell.border = borda

        for r_idx, row in enumerate(
            resumo_anual.itertuples(index=False),
            start=2
        ):

            for c_idx, valor in enumerate(
                row,
                start=1
            ):

                cell = ws_anual.cell(
                    row=r_idx,
                    column=c_idx,
                    value=valor
                )

                cell.border = borda

                if c_idx > 1:

                    cell.number_format = (
                        'R$ #,##0.00'
                    )

        for coluna in range(
            1,
            ws_anual.max_column + 1
        ):

            ws_anual.column_dimensions[
                chr(64 + coluna)
            ].width = 20


        # ----------------------------------------------------
        # GRÁFICO EXCEL
        # ----------------------------------------------------

        chart = LineChart()

        chart.title = (
            "Evolução dos investimentos"
        )

        chart.y_axis.title = "Valor (R$)"
        chart.x_axis.title = "Mês"

        dados = Reference(
            ws_mensal,
            min_col=8,
            max_col=11,
            min_row=1,
            max_row=ws_mensal.max_row
        )

        categorias = Reference(
            ws_mensal,
            min_col=1,
            min_row=2,
            max_row=ws_mensal.max_row
        )

        chart.add_data(
            dados,
            titles_from_data=True
        )

        chart.set_categories(
            categorias
        )

        chart.height = 12
        chart.width = 24

        ws_graficos.add_chart(
            chart,
            "A1"
        )


        # ----------------------------------------------------
        # CONGELAR PAINÉIS
        # ----------------------------------------------------

        ws_mensal.freeze_panes = "A2"
        ws_anual.freeze_panes = "A2"


        wb.save(arquivo)

        arquivo.seek(0)

        return arquivo


    excel = gerar_excel(
        df,
        resumo_anual
    )


    st.download_button(
        label="📥 BAIXAR EXCEL",
        data=excel,
        file_name="VALORES_PROJECAO.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True
    )


# ============================================================
# RODAPÉ
# ============================================================

st.markdown("---")

st.caption(
    "VALORES 6 • Os valores financeiros iniciais são "
    "informados pelo usuário durante a sessão."
)
